import io
import json
import re
import random
from typing import List, Optional

import pandas as pd
from fastapi import APIRouter, UploadFile, File, HTTPException, Depends
from pydantic import BaseModel

from database import get_connection
from auth import require_admin

router = APIRouter()

REQUIRED_COLUMNS = [
    "subject",
    "topic",
    "question",
    "option_a",
    "option_b",
    "option_c",
    "option_d",
    "correct_option",
]

VALID_DIFFICULTIES = {"Easy", "Medium", "Hard"}


# -------------------------------------------------------------------
# Models
# -------------------------------------------------------------------
class ParsedQuestion(BaseModel):
    row_number: int
    subject: str
    topic: str
    question: str
    option_a: str
    option_b: str
    option_c: str
    option_d: str
    correct_option: str
    difficulty: str = "Medium"
    difficulty_source: str = "Auto"
    error: Optional[str] = None


class PreviewResponse(BaseModel):
    total_rows: int
    valid_count: int
    error_count: int
    questions: List[ParsedQuestion]


class DifficultyUpdate(BaseModel):
    difficulty: str


class TopicItem(BaseModel):
    topic_id: int
    topic_name: str
    question_count: int
    easy_count: int = 0
    medium_count: int = 0
    hard_count: int = 0


class SubjectItem(BaseModel):
    subject_id: int
    subject_name: str
    topics: List[TopicItem]


# -------------------------------------------------------------------
# Helpers
# -------------------------------------------------------------------
def clean(value) -> str:
    if value is None:
        return ""
    try:
        if pd.isna(value):
            return ""
    except Exception:
        pass
    return str(value).strip()


def normalize_difficulty(value) -> Optional[str]:
    value = clean(value).lower()
    if value in {"easy", "e", "1"}:
        return "Easy"
    if value in {"medium", "med", "m", "2"}:
        return "Medium"
    if value in {"hard", "h", "3", "difficult"}:
        return "Hard"
    return None


def auto_detect_difficulty(question: str, options: List[str]) -> str:
    """
    Local heuristic only. It does not call an external AI service.
    Excel/file difficulty always takes priority over this detector.
    """
    text = f"{question} {' '.join(options)}".lower()
    score = 0

    hard_patterns = [
        r"\bcalculate\b", r"\bderive\b", r"\bprove\b", r"\banalyze\b",
        r"\bevaluate\b", r"\bdesign\b", r"\bimplement\b", r"\bcompare\b",
        r"\bdifferentiate\b", r"\bjustify\b", r"\boptimize\b",
        r"\bsubnet\b", r"\balgorithm\b", r"\bcomplexity\b",
        r"\barchitecture\b", r"\bscenario\b", r"\bconfiguration\b",
    ]
    medium_patterns = [
        r"\bexplain\b", r"\bdifference\b", r"\bdescribe\b",
        r"\bwhy\b", r"\bhow\b", r"\badvantages?\b", r"\bdisadvantages?\b",
        r"\bfeatures?\b", r"\bfunction\b", r"\bworking\b",
    ]

    score += 2 * sum(bool(re.search(p, question.lower())) for p in hard_patterns)
    score += sum(bool(re.search(p, question.lower())) for p in medium_patterns)

    if len(question.split()) > 22:
        score += 1
    if len(question.split()) > 40:
        score += 1
    if any(len(o.split()) > 25 for o in options):
        score += 1
    if re.search(r"\d+\s*[\+\-\*/=]\s*\d+", text):
        score += 1

    if score >= 3:
        return "Hard"
    if score >= 1:
        return "Medium"
    return "Easy"


def parse_dataframe(df: pd.DataFrame) -> List[ParsedQuestion]:
    # Normalize column names, including BOM/extra spaces.
    df = df.copy()
    df.columns = [str(c).strip().replace("\ufeff", "") for c in df.columns]

    missing = [c for c in REQUIRED_COLUMNS if c not in df.columns]
    if missing:
        raise HTTPException(
            status_code=400,
            detail=f"Missing required columns: {', '.join(missing)}",
        )

    has_difficulty = "difficulty" in df.columns
    parsed = []

    for index, row in df.iterrows():
        row_no = int(index) + 2
        subject = clean(row.get("subject"))
        topic = clean(row.get("topic"))
        question = clean(row.get("question"))
        options = [
            clean(row.get("option_a")),
            clean(row.get("option_b")),
            clean(row.get("option_c")),
            clean(row.get("option_d")),
        ]
        correct = clean(row.get("correct_option")).upper()

        error = None
        for col, value in zip(
            ["subject", "topic", "question", "option_a", "option_b", "option_c", "option_d"],
            [subject, topic, question, *options],
        ):
            if not value:
                error = f"Missing value in '{col}'"
                break

        if not error and correct not in {"A", "B", "C", "D"}:
            error = f"correct_option must be A, B, C, or D (got '{correct}')"

        file_difficulty = normalize_difficulty(row.get("difficulty")) if has_difficulty else None

        if file_difficulty:
            difficulty = file_difficulty
            difficulty_source = "Excel/File"
        else:
            difficulty = auto_detect_difficulty(question, options)
            difficulty_source = "Auto"

        parsed.append(
            ParsedQuestion(
                row_number=row_no,
                subject=subject,
                topic=topic,
                question=question,
                option_a=options[0],
                option_b=options[1],
                option_c=options[2],
                option_d=options[3],
                correct_option=correct,
                difficulty=difficulty,
                difficulty_source=difficulty_source,
                error=error,
            )
        )

    return parsed


def parse_text_questions(text: str) -> List[dict]:
    """
    Supports simple question-paper text formats:
      Q1. Question text
      A. option
      B. option
      C. option
      D. option
      Answer: A
      Difficulty: Easy

    Blank lines separate questions. If difficulty is absent, it is auto detected.
    """
    text = text.replace("\r\n", "\n").replace("\r", "\n").strip()
    blocks = re.split(r"\n\s*\n", text)
    rows = []

    for block in blocks:
        lines = [x.strip() for x in block.split("\n") if x.strip()]
        if not lines:
            continue

        q_match = re.match(r"^(?:Q(?:uestion)?\s*)?(\d+)[\.\):\-]\s*(.*)$", lines[0], re.I)
        question = q_match.group(2).strip() if q_match else lines[0]

        options = {"A": "", "B": "", "C": "", "D": ""}
        answer = ""
        difficulty = ""

        for line in lines[1:]:
            m = re.match(r"^([ABCD])[\.\):\-]\s*(.*)$", line, re.I)
            if m:
                options[m.group(1).upper()] = m.group(2).strip()
                continue

            m = re.match(r"^(?:answer|correct(?:_option)?|correct answer)\s*[:\-]\s*([ABCD])\b", line, re.I)
            if m:
                answer = m.group(1).upper()
                continue

            m = re.match(r"^difficulty\s*[:\-]\s*(easy|medium|hard)\b", line, re.I)
            if m:
                difficulty = normalize_difficulty(m.group(1)) or ""
                continue

        rows.append(
            {
                "subject": "Imported",
                "topic": "General",
                "question": question,
                "option_a": options["A"],
                "option_b": options["B"],
                "option_c": options["C"],
                "option_d": options["D"],
                "correct_option": answer,
                "difficulty": difficulty,
            }
        )

    return rows


def parse_file(filename: str, contents: bytes) -> List[ParsedQuestion]:
    name = (filename or "").lower()
    suffix = name.rsplit(".", 1)[-1] if "." in name else ""

    try:
        if suffix in {"xlsx", "xls"}:
            df = pd.read_excel(io.BytesIO(contents))
            return parse_dataframe(df)

        if suffix == "csv":
            df = pd.read_csv(io.BytesIO(contents))
            return parse_dataframe(df)

        if suffix == "json":
            data = json.loads(contents.decode("utf-8-sig"))
            if isinstance(data, dict):
                data = data.get("questions", data.get("data", [data]))
            df = pd.DataFrame(data)
            return parse_dataframe(df)

        if suffix in {"txt", "md"}:
            rows = parse_text_questions(contents.decode("utf-8-sig", errors="replace"))
            return parse_dataframe(pd.DataFrame(rows))

        if suffix == "docx":
            from docx import Document
            doc = Document(io.BytesIO(contents))
            text = "\n".join(p.text for p in doc.paragraphs if p.text.strip())
            rows = parse_text_questions(text)
            return parse_dataframe(pd.DataFrame(rows))

        if suffix == "pdf":
            from pypdf import PdfReader
            reader = PdfReader(io.BytesIO(contents))
            text = "\n".join(page.extract_text() or "" for page in reader.pages)
            if not text.strip():
                raise HTTPException(
                    status_code=400,
                    detail="This PDF contains no selectable text. Scanned/image PDFs require OCR.",
                )
            rows = parse_text_questions(text)
            return parse_dataframe(pd.DataFrame(rows))

        if suffix == "pptx":
            from pptx import Presentation
            presentation = Presentation(io.BytesIO(contents))
            text_parts = []
            for slide in presentation.slides:
                for shape in slide.shapes:
                    if hasattr(shape, "text") and shape.text.strip():
                        text_parts.append(shape.text)
            rows = parse_text_questions("\n\n".join(text_parts))
            return parse_dataframe(pd.DataFrame(rows))

    except HTTPException:
        raise
    except Exception as exc:
        raise HTTPException(status_code=400, detail=f"Could not read the file: {exc}")

    raise HTTPException(
        status_code=400,
        detail="Unsupported file type. Use .xlsx, .xls, .csv, .json, .txt, .md, .docx, .pdf or .pptx",
    )


# -------------------------------------------------------------------
# Preview / import
# -------------------------------------------------------------------
@router.post("/admin/questions/preview", response_model=PreviewResponse)
async def preview_questions(
    file: UploadFile = File(...),
    admin_user: dict = Depends(require_admin),
):
    contents = await file.read()
    questions = parse_file(file.filename or "", contents)

    valid_count = sum(q.error is None for q in questions)

    return PreviewResponse(
        total_rows=len(questions),
        valid_count=valid_count,
        error_count=len(questions) - valid_count,
        questions=questions,
    )


@router.post("/admin/questions/confirm")
def confirm_questions(
    questions: List[ParsedQuestion],
    admin_user: dict = Depends(require_admin),
):
    conn = get_connection()
    cursor = conn.cursor()
    inserted = 0
    skipped = 0

    try:
        for q in questions:
            if q.error:
                skipped += 1
                continue

            cursor.execute(
                "SELECT subject_id FROM subjects WHERE subject_name = %s",
                (q.subject,),
            )
            subject_row = cursor.fetchone()

            if subject_row:
                subject_id = subject_row["subject_id"]
            else:
                cursor.execute(
                    "INSERT INTO subjects (subject_name) VALUES (%s) RETURNING subject_id",
                    (q.subject,),
                )
                subject_id = cursor.fetchone()["subject_id"]

            cursor.execute(
                """
                SELECT topic_id
                FROM topics
                WHERE subject_id = %s AND topic_name = %s
                """,
                (subject_id, q.topic),
            )
            topic_row = cursor.fetchone()

            if topic_row:
                topic_id = topic_row["topic_id"]
            else:
                cursor.execute(
                    """
                    INSERT INTO topics (subject_id, topic_name)
                    VALUES (%s, %s)
                    RETURNING topic_id
                    """,
                    (subject_id, q.topic),
                )
                topic_id = cursor.fetchone()["topic_id"]

            cursor.execute(
                """
                INSERT INTO questions
                    (
                        subject_id,
                        topic_id,
                        question_text,
                        option_a,
                        option_b,
                        option_c,
                        option_d,
                        correct_option,
                        difficulty
                    )
                VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s)
                """,
                (
                    subject_id,
                    topic_id,
                    q.question,
                    q.option_a,
                    q.option_b,
                    q.option_c,
                    q.option_d,
                    q.correct_option,
                    q.difficulty,
                ),
            )
            inserted += 1

        conn.commit()
        return {"inserted": inserted, "skipped": skipped}

    except Exception:
        conn.rollback()
        raise
    finally:
        cursor.close()
        conn.close()


# -------------------------------------------------------------------
# Question Bank
# -------------------------------------------------------------------
@router.get("/admin/questions")
def list_questions(admin_user: dict = Depends(require_admin)):
    conn = get_connection()
    cursor = conn.cursor()

    try:
        cursor.execute(
            """
            SELECT
                q.question_id,
                q.question_text,
                q.option_a,
                q.option_b,
                q.option_c,
                q.option_d,
                q.correct_option,
                q.difficulty,
                q.subject_id,
                s.subject_name,
                q.topic_id,
                t.topic_name
            FROM questions q
            JOIN subjects s ON s.subject_id = q.subject_id
            JOIN topics t ON t.topic_id = q.topic_id
            ORDER BY q.question_id DESC
            """
        )
        rows = cursor.fetchall()
        return rows
    finally:
        cursor.close()
        conn.close()


@router.patch("/admin/questions/{question_id}/difficulty")
def update_question_difficulty(
    question_id: int,
    request: DifficultyUpdate,
    admin_user: dict = Depends(require_admin),
):
    difficulty = normalize_difficulty(request.difficulty)
    if difficulty not in VALID_DIFFICULTIES:
        raise HTTPException(
            status_code=400,
            detail="difficulty must be Easy, Medium, or Hard",
        )

    conn = get_connection()
    cursor = conn.cursor()

    try:
        cursor.execute(
            "UPDATE questions SET difficulty = %s WHERE question_id = %s",
            (difficulty, question_id),
        )

        if cursor.rowcount == 0:
            conn.rollback()
            raise HTTPException(status_code=404, detail="Question not found")

        conn.commit()
        return {
            "question_id": question_id,
            "difficulty": difficulty,
        }
    finally:
        cursor.close()
        conn.close()


# -------------------------------------------------------------------
# Subjects / topics used by Assign Tests
# -------------------------------------------------------------------
@router.get("/admin/subjects", response_model=List[SubjectItem])
def list_subjects_with_topics(admin_user: dict = Depends(require_admin)):
    conn = get_connection()
    cursor = conn.cursor()

    try:
        cursor.execute(
            "SELECT subject_id, subject_name FROM subjects ORDER BY subject_name"
        )
        subjects = cursor.fetchall()

        result = []

        for subject in subjects:
            cursor.execute(
                """
                SELECT
                    t.topic_id,
                    t.topic_name,
                    COUNT(q.question_id) AS question_count,
                    COUNT(q.question_id) FILTER (WHERE q.difficulty = 'Easy') AS easy_count,
                    COUNT(q.question_id) FILTER (WHERE q.difficulty = 'Medium') AS medium_count,
                    COUNT(q.question_id) FILTER (WHERE q.difficulty = 'Hard') AS hard_count
                FROM topics t
                LEFT JOIN questions q ON q.topic_id = t.topic_id
                WHERE t.subject_id = %s
                GROUP BY t.topic_id, t.topic_name
                ORDER BY t.topic_name
                """,
                (subject["subject_id"],),
            )
            topics = cursor.fetchall()

            result.append(
                SubjectItem(
                    subject_id=subject["subject_id"],
                    subject_name=subject["subject_name"],
                    topics=[TopicItem(**topic) for topic in topics],
                )
            )

        return result
    finally:
        cursor.close()
        conn.close()
